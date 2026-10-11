# -*- coding: utf-8 -*-
"""
生成诊断第6讲和第12讲的举一反三题目文档 (docx)
"""

import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run_font(run, font_name="宋体", font_size_pt=10.5, bold=False, italic=False, color=None):
    """设置 run 的中英文字体、字号与样式"""
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(font_size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_paragraph_with_runs(doc, text_runs, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=3, line_spacing=1.25):
    """
    添加一个段落，支持多个 run 配置
    text_runs: list of tuple (text, font_name, font_size_pt, bold, italic)
    """
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

    for item in text_runs:
        text = item[0]
        font_name = item[1] if len(item) > 1 and item[1] else "宋体"
        size = item[2] if len(item) > 2 and item[2] else 10.5
        bold = item[3] if len(item) > 3 else False
        italic = item[4] if len(item) > 4 else False
        run = p.add_run(text)
        set_run_font(run, font_name=font_name, font_size_pt=size, bold=bold, italic=italic)
    return p


def build_lesson_6():
    """生成第6讲：化学式与化学方程式计算举一反三 docx"""
    doc = Document()
    
    # 页面设置 A4, 边距
    section = doc.sections[0]
    section.page_width = Pt(595.3)
    section.page_height = Pt(841.9)
    section.top_margin = Pt(70.9)
    section.bottom_margin = Pt(70.9)
    section.left_margin = Pt(56.7)
    section.right_margin = Pt(56.7)

    # 1. 大标题
    add_paragraph_with_runs(
        doc,
        [("化学式与化学方程式计算举一反三", "宋体", 15, True, False)],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=6,
        space_after=6,
        line_spacing=1.5
    )

    # 2. 姓名行
    add_paragraph_with_runs(
        doc,
        [("学生姓名：               　。", "宋体", 12, False, False)],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=8,
        line_spacing=1.5
    )

    # 3. 相对原子质量说明
    add_paragraph_with_runs(
        doc,
        [("【相对原子质量：H-1   C-12   N-14   O-16   Na-23   Mg-24   Al-27   S-32   Cl-35.5   K-39   Ca-40   Fe-56   Cu-64   Zn-65】", "宋体", 9.5, False, False)],
        align=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=0,
        space_after=8,
        line_spacing=1.2
    )

    # 考点一
    add_paragraph_with_runs(
        doc,
        [("考点一：化学式的意义与相关计算（第1题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题1
    add_paragraph_with_runs(
        doc,
        [("1、乙酸乙酯（化学式：C₄H₈O₂）是一种具有果香气味的物质，广泛用于食品添加剂中。下列关于乙酸乙酯的说法正确的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．乙酸乙酯由4个碳原子、8个氢原子和2个氧原子构成\n"
          "B．乙酸乙酯的相对分子质量为88 g\n"
          "C．乙酸乙酯中碳、氢、氧三种元素的质量比为4∶8∶2\n"
          "D．乙酸乙酯中碳元素的质量分数计算式为：(12×4) / (12×4 + 1×8 + 16×2) × 100%", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )
    # 题2
    add_paragraph_with_runs(
        doc,
        [("2、布洛芬（化学式为C₁₃H₁₈O₂）具有抗炎、镇痛、解热作用，是常用的退烧药。下列有关布洛芬的说法正确的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．布洛芬中含有碳、氢、氧三种非金属元素\n"
          "B．布洛芬的相对分子质量为206 g\n"
          "C．布洛芬中碳、氢、氧元素的质量比为13∶18∶2\n"
          "D．布洛芬中氢元素的质量分数最大", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )

    # 考点二
    add_paragraph_with_runs(
        doc,
        [("考点二：微观反应示意图与质量守恒定律（第2题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题3
    add_paragraph_with_runs(
        doc,
        [("3、汽车尾气净化装置中发生反应的化学方程式为：2NO + 2CO ==催化剂== N₂ + 2CO₂。下列关于该反应的说法错误的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．该反应中使用催化剂可以增加生成氮气的总质量\n"
          "B．反应涉及的物质中，NO、CO、CO₂均属于氧化物\n"
          "C．反应前后各原子的种类、数目和质量均保持不变\n"
          "D．参加反应的NO与CO的分子个数比为1∶1", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )
    # 题4
    add_paragraph_with_runs(
        doc,
        [("4、科研人员研发出一种新型催化剂，可将甲醛高效转化，反应方程式为：HCHO + O₂ ==催化剂== CO₂ + H₂O。下列说法正确的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．该反应属于置换反应\n"
          "B．该反应前后分子总数发生了改变\n"
          "C．反应消耗的甲醛与氧气的质量比为15∶16\n"
          "D．生成物中的水是由氢原子和氧原子直接构成的", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )

    # 考点三
    add_paragraph_with_runs(
        doc,
        [("考点三：物质微观结构、化学式与方程式基础计算（第3题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题5
    add_paragraph_with_runs(
        doc,
        [("5、乙醇（俗称酒精）是一种常用的液体燃料和消毒剂，每个乙醇分子由2个碳原子、6个氢原子和1个氧原子构成。\n"
          "（1）乙醇的化学式为 _________________；\n"
          "（2）乙醇的相对分子质量为 _______________；\n"
          "（3）乙醇在空气中完全燃烧生成二氧化碳和水，该反应的化学方程式为：\n"
          "     __________________________________________________________________；\n"
          "（4）46 kg 的乙醇完全燃烧，生成的二氧化碳的质量为 _____________ kg。", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=6,
        line_spacing=1.4
    )
    # 题6
    add_paragraph_with_runs(
        doc,
        [("6、天然气的主要成分是甲烷，是一种高热值、低污染的清洁能源。每个甲烷分子由1个碳原子和4个氢原子构成。\n"
          "（1）甲烷的化学式为 _________________；\n"
          "（2）甲烷中碳元素与氢元素的质量比为 _______________；\n"
          "（3）甲烷在氧气中完全燃烧生成二氧化碳和水，该反应的化学方程式为：\n"
          "     __________________________________________________________________；\n"
          "（4）32 g 甲烷完全燃烧，消耗氧气的质量为 _____________ g。", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=6,
        line_spacing=1.4
    )

    # 考点四
    add_paragraph_with_runs(
        doc,
        [("考点四：含杂质物质的化学方程式计算（第4题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题7
    add_paragraph_with_runs(
        doc,
        [("7、赤铁矿是工业炼铁的主要原料，某赤铁矿石样品中氧化铁（Fe₂O₃）的质量分数为80%。工业上用一氧化碳炼铁发生反应：Fe₂O₃ + 3CO ==高温== 2Fe + 3CO₂。理论上用100 t 该赤铁矿石能冶炼出纯铁的质量为 _____________ t。", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=6,
        line_spacing=1.4
    )
    # 题8
    add_paragraph_with_runs(
        doc,
        [("8、辉铜矿石的主要成分为硫化亚铜（Cu₂S）。工业上冶炼铜的一种方法是将辉铜矿在空气中焙烧，反应方程式为：Cu₂S + O₂ ==高温== 2Cu + SO₂。若某批次辉铜矿石中Cu₂S的质量分数为80%，理论上200 t 该辉铜矿石能冶炼出纯铜 _____________ t。", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=6,
        line_spacing=1.4
    )

    # 考点五
    add_paragraph_with_runs(
        doc,
        [("考点五：固体受热分解制气体的综合计算（第5题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题9
    add_paragraph_with_runs(
        doc,
        [("9、实验室常用氯酸钾和二氧化锰的混合物加热制取氧气（反应方程式：2KClO₃ ==MnO₂/△== 2KCl + 3O₂↑）。小明取25.0 g 氯酸钾和二氧化锰的混合物充分加热，完全反应后冷却至室温，称得剩余固体质量为15.4 g。请计算：\n"
          "（1）生成氧气的质量为 _____________ g。\n"
          "（2）原混合物中氯酸钾的质量为多少？（写出规范计算过程）\n\n\n", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=6,
        line_spacing=1.4
    )
    # 题10
    add_paragraph_with_runs(
        doc,
        [("10、石灰石是常见的建筑材料和工业原料。某化学兴趣小组称取20.0 g 石灰石样品（主要成分为CaCO₃，杂质不反应也不受热分解），高温煅烧至质量不再减少，完全反应后冷却并称量，剩余固体质量为13.4 g。（反应方程式：CaCO₃ ==高温== CaO + CO₂↑）请计算：\n"
          "（1）反应生成二氧化碳的质量为 _____________ g。\n"
          "（2）该石灰石样品中碳酸钙的质量分数是多少？（写出规范计算过程）\n\n\n", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=12,
        line_spacing=1.4
    )

    # 参考答案区域
    add_paragraph_with_runs(
        doc,
        [("化学式与化学方程式计算举一反三参考答案", "宋体", 14, True, False)],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=14,
        space_after=8
    )

    # 答案考点一
    add_paragraph_with_runs(doc, [("考点一：化学式的意义与相关计算", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(doc, [("1、D     2、A", "宋体", 10.5, False, False)], space_before=0, space_after=4)

    # 答案考点二
    add_paragraph_with_runs(doc, [("考点二：微观反应示意图与质量守恒定律", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(doc, [("3、A     4、C", "宋体", 10.5, False, False)], space_before=0, space_after=4)

    # 答案考点三
    add_paragraph_with_runs(doc, [("考点三：物质微观结构、化学式与方程式基础计算", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(
        doc,
        [("5、（1）C₂H₅OH（或C₂H₆O）  （2）46  （3）C₂H₅OH + 3O₂ ==点燃== 2CO₂ + 3H₂O  （4）88\n"
          "6、（1）CH₄  （2）3∶1  （3）CH₄ + 2O₂ ==点燃== CO₂ + 2H₂O  （4）128", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=4
    )

    # 答案考点四
    add_paragraph_with_runs(doc, [("考点四：含杂质物质的化学方程式计算", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(
        doc,
        [("7、56\n"
          "8、128", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=4
    )

    # 答案考点五
    add_paragraph_with_runs(doc, [("考点五：固体受热分解制气体的综合计算", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(
        doc,
        [("9、（1）9.6\n"
          "   （2）解：生成氧气的质量 = 25.0 g - 15.4 g = 9.6 g。\n"
          "        设原混合物中氯酸钾的质量为 x。\n"
          "        2KClO₃ ==MnO₂/△== 2KCl + 3O₂↑\n"
          "         245                     96\n"
          "          x                     9.6 g\n"
          "        245 / 96 = x / 9.6 g\n"
          "        x = 24.5 g\n"
          "        答：原混合物中氯酸钾的质量为 24.5 g。\n\n"
          "10、（1）6.6\n"
          "    （2）解：生成二氧化碳的质量 = 20.0 g - 13.4 g = 6.6 g。\n"
          "         设石灰石样品中碳酸钙的质量为 y。\n"
          "         CaCO₃ ==高温== CaO + CO₂↑\n"
          "          100               44\n"
          "           y               6.6 g\n"
          "         100 / 44 = y / 6.6 g\n"
          "         y = 15.0 g\n"
          "         样品中碳酸钙的质量分数 = (15.0 g / 20.0 g) × 100% = 75%\n"
          "         答：该石灰石样品中碳酸钙的质量分数为 75%。", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6
    )

    out_path = os.path.join(os.path.dirname(__file__), "26秋：举一反三第6讲.docx")
    doc.save(out_path)
    print(f"成功生成：{out_path}")


def build_lesson_12():
    """生成第12讲：金属活动性举一反三 docx"""
    doc = Document()
    
    # 页面设置 A4, 边距
    section = doc.sections[0]
    section.page_width = Pt(595.3)
    section.page_height = Pt(841.9)
    section.top_margin = Pt(70.9)
    section.bottom_margin = Pt(70.9)
    section.left_margin = Pt(56.7)
    section.right_margin = Pt(56.7)

    # 1. 大标题
    add_paragraph_with_runs(
        doc,
        [("金属活动性举一反三", "宋体", 15, True, False)],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=6,
        space_after=6,
        line_spacing=1.5
    )

    # 2. 姓名行
    add_paragraph_with_runs(
        doc,
        [("学生姓名：               　。", "宋体", 12, False, False)],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=8,
        line_spacing=1.5
    )

    # 考点一
    add_paragraph_with_runs(
        doc,
        [("考点一：金属与酸、盐溶液反应的判断（第1题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题1
    add_paragraph_with_runs(
        doc,
        [("1、下列金属中，既能与稀盐酸反应生成氢气，又能与硫酸铜溶液反应置换出铜的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．银                B．铁                 C．铜                   D．金", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )
    # 题2
    add_paragraph_with_runs(
        doc,
        [("2、下列金属中，不能与稀硫酸反应，但能与硝酸汞溶液反应置换出金属汞的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．锌                B．铝                 C．铜                   D．银", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )

    # 考点二
    add_paragraph_with_runs(
        doc,
        [("考点二：置换反应的特征与判断（第2题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题3
    add_paragraph_with_runs(
        doc,
        [("3、下列化学反应中，属于置换反应的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．2H₂O₂ ==MnO₂== 2H₂O + O₂↑\n"
          "B．Fe + 2HCl == FeCl₂ + H₂↑\n"
          "C．Ca(OH)₂ + CO₂ == CaCO₃↓ + H₂O\n"
          "D．3CO + Fe₂O₃ ==高温== 2Fe + 3CO₂", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )
    # 题4
    add_paragraph_with_runs(
        doc,
        [("4、下列化学反应中，不属于置换反应的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．Zn + H₂SO₄ == ZnSO₄ + H₂↑\n"
          "B．Cu + 2AgNO₃ == Cu(NO₃)₂ + 2Ag\n"
          "C．CH₄ + 2O₂ ==点燃== CO₂ + 2H₂O\n"
          "D．H₂ + CuO ==加热== Cu + H₂O", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )

    # 考点三
    add_paragraph_with_runs(
        doc,
        [("考点三：金属活动性顺序验证方案评价（第3题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题5
    add_paragraph_with_runs(
        doc,
        [("5、为了验证铝、铁、铜三种金属的活动性顺序，某同学设计的下列实验方案中不可行的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．将铁丝分别插入硫酸铝溶液和硫酸铜溶液中\n"
          "B．将铝丝和铜丝分别插入硫酸亚铁溶液中\n"
          "C．将铁丝和铜丝分别插入稀盐酸中，铝丝插入硫酸亚铁溶液中\n"
          "D．将铝丝分别插入硫酸亚铁溶液和硫酸铜溶液中", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )
    # 题6
    add_paragraph_with_runs(
        doc,
        [("6、为验证锌、铁、银三种金属的活动性强弱，下列各组试剂的选择中不能达到实验目的的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．Zn、Ag 和 FeSO₄ 溶液\n"
          "B．Fe、ZnSO₄ 溶液和 AgNO₃ 溶液\n"
          "C．Zn、Fe、Ag 和稀盐酸\n"
          "D．Zn、FeSO₄ 溶液和 AgNO₃ 溶液", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )

    # 考点四
    add_paragraph_with_runs(
        doc,
        [("考点四：新金属活动性预测与推断（第4题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题7
    add_paragraph_with_runs(
        doc,
        [("7、金属镍（Ni）在材料工业中用途广泛，已知其金属活动性介于铁和铅（Pb）之间。下列对金属镍化学性质的预测中，不合理的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．常温或加热条件下能与氧气反应\n"
          "B．能与稀硫酸反应生成氢气\n"
          "C．能与硫酸铜溶液反应置换出铜\n"
          "D．能与氯化锌溶液反应置换出锌", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )
    # 题8
    add_paragraph_with_runs(
        doc,
        [("8、钛（Ti）被称为“未来金属”，已知钛的金属活动性介于镁和铝之间（Mg > Ti > Al）。下列关于金属钛的化学性质推断错误的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．钛能与稀盐酸反应放出氢气\n"
          "B．钛能置换出硫酸铜溶液中的铜\n"
          "C．钛能置换出硫酸镁溶液中的镁\n"
          "D．钛在一定条件下能与氧气发生反应", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )

    # 考点五
    add_paragraph_with_runs(
        doc,
        [("考点五：金属活动性实验探究综合判断（第5题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题9
    add_paragraph_with_runs(
        doc,
        [("9、某同学为探究铁、铜、银三种金属的活动性顺序，设计了如下实验：①将铁丝插入硫酸铜溶液中；②将铜丝插入硝酸银溶液中；③将铜丝插入稀硫酸中。下列有关说法正确的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．实验①中可观察到铁丝表面析出红色固体，溶液由蓝色逐渐变为黄色\n"
          "B．实验②中发生反应的化学方程式为：Cu + AgNO₃ == CuNO₃ + Ag\n"
          "C．只通过实验①和实验②，就能证明三种金属的活动性顺序为：铁 > 铜 > 银\n"
          "D．实验③中铜丝表面产生大量气泡", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )
    # 题10
    add_paragraph_with_runs(
        doc,
        [("10、某兴趣小组为了探究镁、铁、铜三种金属的活动性顺序，进行了如下实验：\n"
          "实验一：将镁条插入硫酸亚铁溶液中；\n"
          "实验二：将铜片插入硫酸亚铁溶液中。\n"
          "下列说法正确的是（    ）", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=2
    )
    add_paragraph_with_runs(
        doc,
        [("A．实验一中，镁条表面析出红色固体\n"
          "B．实验二中，溶液由浅绿色变为蓝色\n"
          "C．由实验一和实验二可得出三种金属活动性顺序为：Mg > Fe > Cu\n"
          "D．若将实验一中的硫酸亚铁溶液换成硫酸铜溶液，也能得出相同结论", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6,
        line_spacing=1.35
    )

    # 考点六
    add_paragraph_with_runs(
        doc,
        [("考点六：未知金属性质推断与探究实验设计（第6题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题11
    add_paragraph_with_runs(
        doc,
        [("11、金属铟（In）在化合物中常显+3价，是生产液晶显示屏的关键材料。已知铟的金属活动性介于铁和铜之间（Fe > In > Cu）。\n"
          "（1）铟在加热条件下能与氧气反应生成氧化铟（In₂O₃），该反应的化学方程式为：\n"
          "     __________________________________________________________________；\n"
          "（2）铟能与稀盐酸反应生成氯化铟（InCl₃）和氢气，该反应的化学方程式为：\n"
          "     __________________________________________________________________；\n"
          "（3）若要验证Fe、In、Cu三种金属的活动性顺序，现有Fe、Cu单质，还需要选择的一种试剂溶液是 _____________ 溶液。", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=6,
        line_spacing=1.4
    )
    # 题12
    add_paragraph_with_runs(
        doc,
        [("12、金属铬（Cr）在不锈钢和电镀中广泛应用。已知铬在化合物中常显+2价，其金属活动性介于镁和锌之间（Mg > Cr > Zn）。\n"
          "（1）铬在加热条件下与氧气反应生成氧化铬（CrO），该反应的化学方程式为：\n"
          "     __________________________________________________________________；\n"
          "（2）将金属铬放入稀硫酸中，反应生成硫酸铬（CrSO₄）和氢气，该反应的化学方程式为：\n"
          "     __________________________________________________________________；\n"
          "（3）若只用两种试剂通过一步反应就能验证Cr和Zn的活动性强弱，可选用的组合是金属铬和 _____________ 溶液。", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=6,
        line_spacing=1.4
    )

    # 考点七
    add_paragraph_with_runs(
        doc,
        [("考点七：工业流程中金属的回收与分离（第7题）", "宋体", 11, True, False)],
        space_before=6,
        space_after=4
    )
    # 题13
    add_paragraph_with_runs(
        doc,
        [("13、电子工业废料中常含有铜、铁、金（Au）等金属，某工厂设计了如下从废料中回收金属铜和金的工艺流程：\n"
          "废料粉末 ──加入足量稀盐酸──> 滤液A + 滤渣B（含Cu、Au）\n"
          "滤渣B ──在空气中充分灼烧──> 固体C（含CuO、Au） ──加入足量稀盐酸──> 滤液D + 金属金\n"
          "（1）上述流程中多次涉及将固体与液体分离的实验操作，该操作名称是 _____________。\n"
          "（2）向废料中加入足量稀盐酸时，铁发生反应的化学方程式为：\n"
          "     __________________________________________________________________。\n"
          "（3）固体C中加入足量稀盐酸后，氧化铜与稀盐酸反应的化学方程式为：\n"
          "     __________________________________________________________________。", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=6,
        line_spacing=1.4
    )
    # 题14
    add_paragraph_with_runs(
        doc,
        [("14、某工厂排放的废水中含有硫酸亚铁、硫酸铜和硫酸银三种溶质，为回收金属铜、银并得到硫酸亚铁晶体，设计了如下回收方案：\n"
          "废水 ──加入过量金属X──> 过滤得到滤液1和滤渣1\n"
          "滤渣1 ──加入适量稀硫酸──> 过滤得到硫酸亚铁溶液和滤渣2（含Cu、Ag）\n"
          "滤渣2 ──加适量试剂并提纯──> 分离出纯净的铜和银\n"
          "（1）加入的金属X是 ___________（填化学式）。\n"
          "（2）加入过量金属X时，置换出铜的化学反应方程式为：\n"
          "     __________________________________________________________________。\n"
          "（3）在滤渣1中加入稀硫酸的目的是：\n"
          "     __________________________________________________________________。", "宋体", 10.5, False, False)],
        space_before=2,
        space_after=12,
        line_spacing=1.4
    )

    # 参考答案区域
    add_paragraph_with_runs(
        doc,
        [("金属活动性举一反三参考答案", "宋体", 14, True, False)],
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=14,
        space_after=8
    )

    # 答案考点一
    add_paragraph_with_runs(doc, [("考点一：金属与酸、盐溶液反应的判断", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(doc, [("1、B     2、C", "宋体", 10.5, False, False)], space_before=0, space_after=4)

    # 答案考点二
    add_paragraph_with_runs(doc, [("考点二：置换反应的特征与判断", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(doc, [("3、B     4、C", "宋体", 10.5, False, False)], space_before=0, space_after=4)

    # 答案考点三
    add_paragraph_with_runs(doc, [("考点三：金属活动性顺序验证方案评价", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(doc, [("5、D     6、D", "宋体", 10.5, False, False)], space_before=0, space_after=4)

    # 答案考点四
    add_paragraph_with_runs(doc, [("考点四：新金属活动性预测与推断", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(doc, [("7、D     8、C", "宋体", 10.5, False, False)], space_before=0, space_after=4)

    # 答案考点五
    add_paragraph_with_runs(doc, [("考点五：金属活动性实验探究综合判断", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(doc, [("9、C     10、C", "宋体", 10.5, False, False)], space_before=0, space_after=4)

    # 答案考点六
    add_paragraph_with_runs(doc, [("考点六：未知金属性质推断与探究实验设计", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(
        doc,
        [("11、（1）4In + 3O₂ ==加热== 2In₂O₃\n"
          "    （2）2In + 6HCl == 2InCl₃ + 3H₂↑\n"
          "    （3）InCl₃（或硫酸铟 / 硝酸铟）\n"
          "12、（1）2Cr + O₂ ==加热== 2CrO\n"
          "    （2）Cr + H₂SO₄ == CrSO₄ + H₂↑\n"
          "    （3）硫酸锌（或ZnSO₄ / ZnCl₂）", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=4
    )

    # 答案考点七
    add_paragraph_with_runs(doc, [("考点七：工业流程中金属的回收与分离", "宋体", 10.5, True, False)], space_before=4, space_after=2)
    add_paragraph_with_runs(
        doc,
        [("13、（1）过滤\n"
          "    （2）Fe + 2HCl == FeCl₂ + H₂↑\n"
          "    （3）CuO + 2HCl == CuCl₂ + H₂O\n"
          "14、（1）Fe\n"
          "    （2）Fe + CuSO₄ == FeSO₄ + Cu\n"
          "    （3）除去过量的铁粉（或使过量的铁完全转化为硫酸亚铁）", "宋体", 10.5, False, False)],
        space_before=0,
        space_after=6
    )

    out_path = os.path.join(os.path.dirname(__file__), "26秋：举一反三第12讲.docx")
    doc.save(out_path)
    print(f"成功生成：{out_path}")


if __name__ == "__main__":
    print("开始生成练习题文档...")
    build_lesson_6()
    build_lesson_12()
    print("全部文档生成完毕！")
